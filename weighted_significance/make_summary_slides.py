# Summary slides (pdf + pptx): per-channel phiCP gain over the run-3 method, and combined totals.
# Needs python-pptx (installed under claude_scratch_minitest/pylib: PYTHONPATH=<that>/pylib python make_summary_slides.py).
"""Two slides (pptx + pdf) with the per-channel phiCP gains vs the run-3 method."""
import sys, glob, os, numpy as np
sys.path.insert(0, '/vols/cms/dw515/DiTauEntanglement_new/TauPolaris/weighted_significance')
from total_asymmetry import load_asymmetry_yaml, compute_total_asymmetry
WS = '/vols/cms/dw515/DiTauEntanglement_new/TauPolaris/weighted_significance/'
run3 = {}
for f in glob.glob('/vols/cms/dw515/DiTauEntanglement_new/TauPolaris/outputs_Flow_Uncorr_Masked_Hadronic_100e_July28/phiCP/logs/*_RecoRun3.npz'):
    z = np.load(f); e, o = z['even_counts'], z['odd_counts']; run3[os.path.basename(f).split('_')[0]] = float(np.sqrt(((e-o)**2).sum())/e.sum())
methods = {'nu-method': load_asymmetry_yaml(WS+'asymmetry_had_Polaris.yaml'), 'flow Sep01': load_asymmetry_yaml(WS+'asymmetry_had_FlowSep01.yaml'),
           'transformer': load_asymmetry_yaml(WS+'asymmetry_had_Regressor_Sep25.yaml'),
           'vis+SV circ.': load_asymmetry_yaml(WS+'asymmetry_had_FlowVisSVSep22_circmean100.yaml'),
           'vis+all circ.': load_asymmetry_yaml(WS+'asymmetry_had_FlowVisAllSep22_circmean100.yaml')}
MAIN = list(methods)   # 'best of all' picks the best column per channel
lep = load_asymmetry_yaml(WS+'asymmetry_lep_Polaris.yaml')
order = ['DM0DM0','DM0DM1','DM1DM1','DM0DM2','DM1DM2','DM2DM2','DM0DM10','DM1DM10','DM2DM10','DM10DM10']
dm11 = ['DM0DM11','DM1DM11','DM2DM11','DM10DM11','DM11DM11']
lab = {'0':'1π0π⁰','1':'1π1π⁰','2':'1π2π⁰','10':'3π0π⁰','11':'3π1π⁰'}
name = lambda k: ' – '.join(lab[x] for x in k.replace('DM','',1).split('DM'))
pct = lambda v, r: ('−' if v < r else '+') + f'{abs(100*(v/r-1)):.1f}%'
N = lambda k: methods['nu-method'][k]['N']

# ---- content
T1_TITLE = 'φCP sensitivity per channel: gain over the run-3 method'
T1_SUB = 'CP-even vs CP-odd quadrature asymmetry of φCP, H→ττ, Jul20 sample (3.44M events), same events for every method. Leptonic channels not shown.'
rows1 = [['channel (τ⁺ – τ⁻)', 'N', 'run-3 asym.', 'nu-method', 'flow Sep01', 'transformer', 'vis+SV circ.', 'vis+all circ.', 'best of all']]
for k in order:
    r = run3[k]; vals = [d[k]['asymmetry'] for d in methods.values()]
    rows1.append([name(k), f'{N(k):,}', f'{r:.4f}'] + [pct(v, r) for v in vals] + [pct(max(methods[m][k]['asymmetry'] for m in MAIN), r)])
t3 = np.sqrt(sum(N(k)*run3[k]**2 for k in order)/sum(N(k) for k in order))
sub = lambda d: {k: d[k] for k in order}
best10 = {k: max((methods[m][k] for m in MAIN), key=lambda e: e['asymmetry']) for k in order}
rows1.append(['total, these 10 channels', '', f'{t3:.4f}'] + [pct(compute_total_asymmetry(sub(d)), t3) for d in methods.values()] + [pct(compute_total_asymmetry(best10), t3)])
W1 = [2.2, 1.1, 1.15, 1.25, 1.25, 1.3, 1.35, 1.35, 1.3]
NOTES1 = ['vis+SV / vis+all = Sep01 flow + visible-τ token + secondary vertex (+ all per-leg vectors) in the leg’s (n,r,k) frame, 100 epochs;',
          '“circ.” = φCP from the circular mean of 100 flow samples per event. Flow Sep01 uses the MAP. MAP results for vis, vis+SV, vis+all: separate slide.',
          'Run-3 method = classic acoplanarity from impact parameters / decay planes. Totals are N-weighted: √(Σ N·a² / Σ N).',
          '"best of all" picks the best column per channel; vis+all adds +5–17% over vis+SV in the 3π1π⁰ channels only (see slide 2).']
T2_TITLE = 'Channels with a 3π1π⁰ leg, and combined totals'
T2_SUB = 'Run-3 has no 3π1π⁰ branch, so these channels are shown as asymmetries. The combined totals include the leptonic channels (nu-method values for every row).'
rows2 = [['channel (τ⁺ – τ⁻)', 'N', 'nu-method', 'flow Sep01', 'transformer', 'vis+SV circ.', 'vis+all circ.']] + \
        [[name(k), f'{N(k):,}'] + [f"{d[k]['asymmetry']:.4f}" for d in methods.values()] for k in dm11]
W2 = [1.8, 0.85, 1.0, 1.05, 1.1, 1.1, 1.1]
ref = compute_total_asymmetry({**methods['nu-method'], **lep}); old = ref / 1.18
allbest = {k: max((methods[m][k] for m in MAIN), key=lambda e: e['asymmetry']) for k in methods['nu-method']}
rows3 = [['all channels', 'total asym.', 'vs old method']] + \
        [[nm, f'{compute_total_asymmetry({**d, **lep}):.4f}', pct(compute_total_asymmetry({**d, **lep}), old)] for nm, d in list(methods.items()) + [('best per channel', allbest)]]
W3 = [1.55, 1.0, 1.25]
NOTES2 = [          'Old-method total (0.0529) inferred from the', 'published +18% of the nu-method; every', 'other row uses the same denominator.', '',
          'Transformer = 25-epoch multi-target', 'regressor (pol-vecs, neutrinos, taus) with', 'the flow’s conditioner architecture.']
NAVY, INK, MUTED, ROW, GOOD, BAD = '1E2761', '212121', '6B7280', 'EEF2F8', '1B6E3A', 'B02A2A'

# ---- pdf (matplotlib)
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
def draw_table(ax, rows, col_w, x0, y0, rh=0.34, size=11, best_col=None, pct_cols=(), bold_last=False):
    for i, row in enumerate(rows):
        y = y0 + i*rh; x = x0
        for j, val in enumerate(row):
            w = col_w[j]; fc = '#'+NAVY if i == 0 else ('#'+ROW if i % 2 == 0 else 'white')
            ax.add_patch(plt.Rectangle((x, y), w, rh, facecolor=fc, edgecolor='white', linewidth=1))
            color = 'white' if i == 0 else '#'+INK; bold = i == 0 or (best_col is not None and j == best_col) or (bold_last and i == len(rows)-1)
            if i > 0 and j in pct_cols and str(val).endswith('%'): color = '#'+BAD if str(val).startswith('−') else '#'+GOOD
            ax.text(x + 0.08 if j == 0 else x + w - 0.08, y + rh/2, str(val), ha='left' if j == 0 else 'right', va='center', fontsize=size, color=color, fontweight='bold' if bold else 'normal')
            x += w
def page(title, subtitle):
    fig = plt.figure(figsize=(13.333, 7.5)); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 13.333); ax.set_ylim(7.5, 0); ax.axis('off')
    ax.text(0.6, 0.65, title, fontsize=24, fontweight='bold', color='#'+NAVY, va='center'); ax.text(0.6, 1.15, subtitle, fontsize=10.5, color='#'+MUTED, va='center'); return fig, ax
with PdfPages(WS + 'phiCP_gains_vs_run3_Sep29.pdf') as pdf:
    fig, ax = page(T1_TITLE, T1_SUB); draw_table(ax, rows1, W1, 0.6, 1.5, best_col=8, pct_cols=(3,4,5,6,7,8), bold_last=True, size=10)
    for i, l in enumerate(NOTES1): ax.text(0.6, 5.85 + 0.3*i, l, fontsize=9.5, color='#'+MUTED, wrap=True)
    pdf.savefig(fig); plt.close(fig)
    fig, ax = page(T2_TITLE, T2_SUB); draw_table(ax, rows2, W2, 0.4, 1.5, size=10); draw_table(ax, rows3, W3, 9.1, 1.5, pct_cols=(2,), size=9.5)
    for i, l in enumerate(NOTES2): ax.text(9.1, 4.6 + 0.25*i, l, fontsize=9, color='#'+MUTED)
    pdf.savefig(fig); plt.close(fig)

# ---- pptx (python-pptx)
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
rgb = lambda h: RGBColor.from_string(h)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5); blank = prs.slide_layouts[6]
def text(slide, x, y, w, h, s, size=14, bold=False, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run(); r.text = s; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = 'Calibri'
def ptable(slide, x, y, rows, col_w, size=12, best_col=None, pct_cols=(), bold_last=False):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(sum(col_w)), Inches(0.34*len(rows))); t = shp.table
    for j, cw in enumerate(col_w): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(0.34)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.02); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = 'Calibri'; c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = rgb(NAVY); r.font.bold = True; r.font.color.rgb = rgb('FFFFFF')
            else:
                c.fill.fore_color.rgb = rgb(ROW) if i % 2 == 0 else rgb('FFFFFF'); r.font.color.rgb = rgb(INK)
                if j in pct_cols and str(val).endswith('%'): r.font.color.rgb = rgb(BAD) if str(val).startswith('−') else rgb(GOOD)
                if (best_col is not None and j == best_col) or (bold_last and i == len(rows)-1): r.font.bold = True
s = prs.slides.add_slide(blank); text(s, 0.6, 0.35, 12.1, 0.6, T1_TITLE, 30, True, NAVY); text(s, 0.6, 0.95, 12.1, 0.4, T1_SUB, 12, False, MUTED)
ptable(s, 0.6, 1.5, rows1, W1, size=10, best_col=8, pct_cols=(3,4,5,6,7,8), bold_last=True); text(s, 0.6, 5.75, 12.1, 1.3, ' '.join(NOTES1), 10, False, MUTED)
s = prs.slides.add_slide(blank); text(s, 0.6, 0.35, 12.1, 0.6, T2_TITLE, 30, True, NAVY); text(s, 0.6, 0.95, 12.1, 0.4, T2_SUB, 12, False, MUTED)
ptable(s, 0.4, 1.5, rows2, W2, size=10); ptable(s, 9.1, 1.5, rows3, W3, size=9, pct_cols=(2,)); text(s, 9.1, 4.6, 3.8, 2.4, ' '.join(l for l in NOTES2 if l), 11, False, MUTED)
prs.save(WS + 'phiCP_gains_vs_run3_Sep29.pptx'); print('written pptx + pdf')
