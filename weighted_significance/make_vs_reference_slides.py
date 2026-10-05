"""One slide (pdf + pptx) comparing one or more models with a reference model, channel by channel.

    python weighted_significance/make_vs_reference_slides.py --ref asymmetry_had_FlowSep01.yaml --ref_label "flow Sep01" \
        --model asymmetry_had_FlowVisSep22.yaml "flow vis" --out phiCP_vis_vs_Sep01
"""
import argparse, os, sys
import numpy as np
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from total_asymmetry import load_asymmetry_yaml, compute_total_asymmetry

ORDER = ['DM0DM0', 'DM0DM1', 'DM1DM1', 'DM0DM2', 'DM1DM2', 'DM2DM2', 'DM0DM10', 'DM1DM10', 'DM2DM10', 'DM10DM10',
         'DM0DM11', 'DM1DM11', 'DM2DM11', 'DM10DM11', 'DM11DM11']
LAB = {'0': '1π0π⁰', '1': '1π1π⁰', '2': '1π2π⁰', '10': '3π0π⁰', '11': '3π1π⁰'}
name = lambda k: ' – '.join(LAB[x] for x in k.replace('DM', '', 1).split('DM'))
NAVY, INK, MUTED, ROW, GOOD, BAD = '1E2761', '212121', '6B7280', 'EEF2F8', '1B6E3A', 'B02A2A'
sgn = lambda v: ('−' if v < 0 else '+') + f'{abs(100 * v):.1f}%'

ap = argparse.ArgumentParser()
ap.add_argument('--ref', required=True); ap.add_argument('--ref_label', required=True)
ap.add_argument('--model', nargs=2, action='append', required=True, metavar=('YAML', 'LABEL'))
ap.add_argument('--lep', default='asymmetry_lep_Polaris.yaml'); ap.add_argument('--nu_ref', default='asymmetry_had_Polaris.yaml')
ap.add_argument('--loss_dir', nargs=2, action='append', default=[], metavar=('LABEL', 'MODEL_DIR'),
                help='add the training/validation loss for the column LABEL, read from MODEL_DIR/plots/partial_meta.pth')
ap.add_argument('--out', required=True); ap.add_argument('--title', default=None); ap.add_argument('--note', default='')
a = ap.parse_args()
P = lambda f: f if os.path.isabs(f) else os.path.join(HERE, f)
ref = load_asymmetry_yaml(P(a.ref)); models = [(lab, load_asymmetry_yaml(P(f))) for f, lab in a.model]
lep = load_asymmetry_yaml(P(a.lep)); old = compute_total_asymmetry({**load_asymmetry_yaml(P(a.nu_ref)), **lep}) / 1.18
for lab, d in models:
    # same events: counts may differ only by the odd event whose estimate is undefined (e.g. all flow samples failed)
    assert all(abs(d[k]['N'] - ref[k]['N']) <= max(2, 1e-4 * ref[k]['N']) for k in ORDER), f'{lab}: event counts differ from the reference -- not the same events'

hdr = ['channel (τ⁺ – τ⁻)', 'N', a.ref_label] + sum([[lab, 'vs ' + a.ref_label] for lab, _ in models], [])
rows = [hdr]
for k in ORDER:
    r = ref[k]['asymmetry']; row = [name(k), f"{ref[k]['N']:,}", f'{r:.4f}']
    for lab, d in models: row += [f"{d[k]['asymmetry']:.4f}", sgn(d[k]['asymmetry'] / r - 1)]
    rows.append(row)
tot = lambda d: compute_total_asymmetry({k: d[k] for k in ORDER})
rows.append(['total, hadronic (15 ch.)', '', f'{tot(ref):.4f}'] + sum([[f'{tot(d):.4f}', sgn(tot(d) / tot(ref) - 1)] for _, d in models], []))
comb = lambda d: compute_total_asymmetry({**{k: d[k] for k in ORDER}, **lep})
rows.append(['total, had + lep', '', f'{comb(ref):.4f}'] + sum([[f'{comb(d):.4f}', sgn(comb(d) / comb(ref) - 1)] for _, d in models], []))
rows.append(['   vs old method', '', sgn(comb(ref) / old - 1)] + sum([[sgn(comb(d) / old - 1), ''] for _, d in models], []))
# ---- loss rows (validation NLL at the best epoch, training NLL at that epoch). Only comparable when
# the models share targets, data and output normalisation -- checked here when both dirs are given.
loss = {}
if a.loss_dir:
    import torch
    norms = {}
    for lab, d in a.loss_dir:
        h = torch.load(os.path.join(d, 'plots', 'partial_meta.pth'), map_location='cpu')['history']
        va, tr = np.array(h['val_loss']), np.array(h['train_loss']); b = int(np.argmin(va))
        loss[lab] = (va[b], tr[b], b + 1, len(va))
        norms[lab] = np.load(os.path.join(d, 'normalization_params.npz'))
    labs = list(norms)
    for l2 in labs[1:]:
        for key in ('output_mean', 'output_std'):
            assert np.allclose(norms[labs[0]][key], norms[l2][key]), f'{l2}: different output normalisation, losses not comparable'
    cols = [a.ref_label] + [lab for lab, _ in models]
    def lrow(title, idx, fmt):
        r = [title, '', fmt(loss[a.ref_label][idx]) if a.ref_label in loss else '']
        for lab, _ in models:
            if lab in loss:
                diff = '' if a.ref_label not in loss or idx >= 2 else f'{loss[lab][idx] - loss[a.ref_label][idx]:+.3f}'
                r += [fmt(loss[lab][idx]), diff]
            else:
                r += ['', '']
        return r
    rows.append(lrow('best val. loss (NLL)', 0, lambda v: f'{v:.3f}'))
    rows.append(lrow('train loss, same epoch', 1, lambda v: f'{v:.3f}'))
    rows.append(lrow('best epoch / epochs', 2, lambda v: ''))
    for i, lab in enumerate([a.ref_label] + [l for l, _ in models]):
        if lab in loss:
            col = 2 if i == 0 else 3 + 2 * (i - 1)
            rows[-1][col] = f'{loss[lab][2]} / {loss[lab][3]}'
n_tot = 3
title = a.title or ' / '.join(lab for lab, _ in models) + f' vs {a.ref_label}: φCP sensitivity per channel'
sub = (f'CP-even vs CP-odd quadrature asymmetry of φCP (20 bins), H→ττ Jul20 sample, identical events and weights for every column. '
       f'Leptonic channels use the nu-method for every column.')
W = [2.55, 1.15, 1.35] + [1.35, 1.45] * len(models)
if sum(W) > 12.15:                                   # many models: scale the columns to fit the slide
    W = [w * 12.15 / sum(W) for w in W]
pct_cols = {4 + 2 * i for i in range(len(models))}
notes = [a.note] if a.note else []
notes += (['Loss = flow negative log-likelihood per event on the normalised targets (lower is better); identical output normalisation verified.'] if loss else [])
notes += ['Totals are N-weighted: √(Σ N·a² / Σ N). "vs old method" uses the old-method total (0.0529) inferred from the published nu-method +18%.']

import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
fig = plt.figure(figsize=(13.333, 7.5)); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 13.333); ax.set_ylim(7.5, 0); ax.axis('off')
ax.text(0.6, 0.55, title, fontsize=22, fontweight='bold', color='#' + NAVY, va='center')
ax.text(0.6, 0.98, sub, fontsize=9.5, color='#' + MUTED, va='center')
note_room = 13.333 - 0.5 - (0.6 + sum(W) + 0.4)
# rows shrink when the notes have to go underneath the table (no room beside it)
rh = 0.27 if note_room > 3.0 else min(0.27, (7.5 - 1.25 - 1.3) / len(rows))
fs = (10 if rh >= 0.25 else 9) - (1 if len(models) >= 3 else 0)
for i, row in enumerate(rows):
    y = 1.25 + i * rh; x = 0.6; is_tot = len(ORDER) < i <= len(ORDER) + 3
    for j, val in enumerate(row):
        w = W[j]; fc = '#' + NAVY if i == 0 else ('#' + ROW if i % 2 == 0 else 'white')
        ax.add_patch(plt.Rectangle((x, y), w, rh, facecolor=fc, edgecolor='white', lw=1))
        col = 'white' if i == 0 else '#' + INK
        if i > 0 and (j in pct_cols or (is_tot and str(val).endswith('%'))) and str(val).endswith('%'):
            col = '#' + (BAD if str(val).startswith('−') else GOOD)
        ax.text(x + 0.08 if j == 0 else x + w - 0.08, y + rh / 2, str(val), ha='left' if j == 0 else 'right', va='center',
                fontsize=fs, color=col, fontweight='bold' if (i == 0 or is_tot) else 'normal')
        x += w
import textwrap
note_x = 0.6 + sum(W) + 0.4; note_w = 13.333 - 0.5 - note_x
side = note_w > 3.0            # notes to the right of the table when there is room, else below it
nx, ny = (note_x, 1.35) if side else (0.6, 1.25 + len(rows) * rh + 0.3)
width_chars = int((note_w if side else 12.0) * 13)
y = ny
for l in notes:
    for line in textwrap.wrap(l, width_chars):
        ax.text(nx, y, line, fontsize=9, color='#' + MUTED, va='top'); y += 0.2
    y += 0.1
with PdfPages(os.path.join(HERE, a.out + '.pdf')) as pdf:
    pdf.savefig(fig)
plt.close(fig)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
rgb = RGBColor.from_string
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5); s = prs.slides.add_slide(prs.slide_layouts[6])
def text(x, y, w, h, t, size, bold=False, color=INK):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = 'Calibri'
text(0.6, 0.3, 12.1, 0.5, title, 26, True, NAVY); text(0.6, 0.85, 12.1, 0.35, sub, 11, False, MUTED)
tb = s.shapes.add_table(len(rows), len(hdr), Inches(0.6), Inches(1.25), Inches(sum(W)), Inches(rh * len(rows))).table
for j, w in enumerate(W): tb.columns[j].width = Inches(w)
for i, row in enumerate(rows):
    tb.rows[i].height = Inches(rh); is_tot = len(ORDER) < i <= len(ORDER) + 3
    for j, val in enumerate(row):
        c = tb.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.01); c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        r = p.add_run(); r.text = str(val); r.font.size = Pt(fs + 1); r.font.name = 'Calibri'; c.fill.solid()
        c.fill.fore_color.rgb = rgb(NAVY) if i == 0 else (rgb(ROW) if i % 2 == 0 else rgb('FFFFFF'))
        r.font.color.rgb = rgb('FFFFFF') if i == 0 else rgb(INK); r.font.bold = i == 0 or is_tot
        if i > 0 and str(val).endswith('%') and (j in pct_cols or is_tot):
            r.font.color.rgb = rgb(BAD) if str(val).startswith('−') else rgb(GOOD)
text(nx, ny, (note_w if side else 12.1), 3.0, '\n\n'.join(notes), 10, False, MUTED)
prs.save(os.path.join(HERE, a.out + '.pptx'))
print('written', a.out + '.pdf/.pptx')
