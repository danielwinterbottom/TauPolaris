# Summary slides (pdf + pptx), Oct 06: per-channel phiCP gain over the run-3 method for the nu-method,
# the vis+all transformer regressor (hadronic AND semileptonic), the Sep01 flow and the 25-epoch
# pretrained-conditioner flow (+64 layer, unfrozen) -- the flows are hadronic only.
# Needs python-pptx: PYTHONPATH=/vols/cms/dw515/DiTauEntanglement_new/claude_scratch_minitest/pylib python make_summary_slides_Oct06.py
import sys, glob, os, numpy as np
WS = os.path.dirname(os.path.abspath(__file__)) + '/'
sys.path.insert(0, WS)
from total_asymmetry import load_asymmetry_yaml, compute_total_asymmetry as T
OUT = WS + 'phiCP_gains_vs_run3_Oct06'

run3 = {}
for f in glob.glob(WS + '../outputs_Flow_Uncorr_Masked_Hadronic_100e_July28/phiCP/logs/*_RecoRun3.npz'):
    z = np.load(f); e, o = z['even_counts'], z['odd_counts']; run3[os.path.basename(f).split('_')[0]] = float(np.sqrt(((e-o)**2).sum())/e.sum())
# semileptonic run-3 ("Approximate") values: TauPolaris paper table (no local recoRun3 histograms for these channels)
run3.update({'DM100DM0': 0.041, 'DM100DM1': 0.044, 'DM100DM2': 0.032, 'DM100DM10': 0.036})
Y = lambda f: load_asymmetry_yaml(WS + f)
had = {'nu-method': Y('asymmetry_had_Polaris.yaml'), 'transformer': Y('asymmetry_had_RegressorVisAll.yaml'),
       'flow Sep01 (MAP)': Y('asymmetry_had_FlowSep01.yaml'), '25e +64 (MAP)': Y('asymmetry_had_PretrainedCondHidden25e.yaml'),
       '25e +64 (circ.)': Y('asymmetry_had_PretrainedCondHidden25e_circmean100.yaml')}
lep = {'nu-method': Y('asymmetry_lep_Polaris.yaml'), 'transformer': Y('asymmetry_lep_RegressorSemilepVisAll.yaml')}
nu_had = had['nu-method']
for m, d in list(had.items()) + list(lep.items()):
    ref = nu_had if 'DM0DM0' in d else lep['nu-method']
    assert all(abs(d[k]['N'] - ref[k]['N']) <= 2 for k in ref), f'{m}: not the same events'

order = ['DM0DM0', 'DM0DM1', 'DM1DM1', 'DM0DM2', 'DM1DM2', 'DM2DM2', 'DM0DM10', 'DM1DM10', 'DM2DM10', 'DM10DM10']
dm11 = ['DM0DM11', 'DM1DM11', 'DM2DM11', 'DM10DM11', 'DM11DM11']
lep4 = ['DM100DM0', 'DM100DM1', 'DM100DM2', 'DM100DM10']
lab = {'0': '1π0π⁰', '1': '1π1π⁰', '2': '1π2π⁰', '10': '3π0π⁰', '11': '3π1π⁰', '100': 'τℓ'}
name = lambda k: ' – '.join(lab[x] for x in k.replace('DM', '', 1).split('DM'))
pct = lambda v, r: ('−' if v < r else '+') + f'{abs(100*(v/r-1)):.1f}%'
a = lambda d, k: d[k]['asymmetry']

def tot(d, keys, src=None):
    src = src or d
    N = np.array([src[k]['N'] for k in keys], float); x = np.array([d[k]['asymmetry'] if isinstance(d[k], dict) else d[k] for k in keys])
    return np.sqrt((N * x**2).sum() / N.sum())

# ---- slide 1: hadronic channels with a run-3 value
T1 = 'φCP sensitivity per hadronic channel: gain over the run-3 method'
S1 = 'CP-even vs CP-odd quadrature asymmetry of φCP (20 bins), H→ττ Jul20 sample, identical events for every method.'
rows1 = [['channel (τ⁺ – τ⁻)', 'N', 'run-3 asym.'] + list(had)]
for k in order:
    rows1.append([name(k), f"{nu_had[k]['N']:,}", f'{run3[k]:.4f}'] + [pct(a(d, k), run3[k]) for d in had.values()])
t3 = tot(run3, order, nu_had)
rows1.append(['total, these 10 channels', '', f'{t3:.4f}'] + [pct(tot(d, order), t3) for d in had.values()])
W1 = [2.2, 1.0, 1.15, 1.35, 1.4, 1.65, 1.5, 1.55]
N1 = ['transformer = vis+all transformer regressor (pol-vecs, neutrinos, taus; frame-aligned inputs, Oct04). flow Sep01 = 100-epoch flow, 72 inputs.',
      '25e +64 = 25-epoch flow, conditioner initialised from the 72-input regressor incl. its 64→64 layer, unfrozen. MAP = gradient-search mode;',
      'circ. = circular mean of 100 flow samples per event (not available for Sep01 on the full sample). No uncertainty cut applied.',
      'Run-3 method = classic acoplanarity from impact parameters / decay planes. Totals are N-weighted: √(Σ N·a² / Σ N).']

# ---- slide 2: 3pi1pi0 channels + semileptonic channels
T2 = '3π1π⁰ channels and semileptonic channels'
S2 = 'Run-3 has no 3π1π⁰ branch, so those channels are shown as asymmetries. Semileptonic: run-3 values from the TauPolaris paper table.'
rows2 = [['channel (τ⁺ – τ⁻)', 'N'] + list(had)] + [[name(k), f"{nu_had[k]['N']:,}"] + [f'{a(d, k):.4f}' for d in had.values()] for k in dm11]
W2 = [1.75, 0.85, 1.1, 1.2, 1.6, 1.45, 1.5]
lep_nu, lep_T = lep['nu-method'], lep['transformer']
rows3 = [['channel', 'N', 'run-3', 'nu-method', 'transformer']]
for k in lep4:
    rows3.append([name(k), f"{lep_nu[k]['N']:,}", f'{run3[k]:.3f}', pct(a(lep_nu, k), run3[k]), pct(a(lep_T, k), run3[k])])
rows3.append([name('DM100DM11'), f"{lep_nu['DM100DM11']['N']:,}", '–', f"{a(lep_nu, 'DM100DM11'):.4f}", f"{a(lep_T, 'DM100DM11'):.4f}"])
r3l = tot(run3, lep4, lep_nu)
rows3.append(['total, 4 channels', '', f'{r3l:.4f}', pct(tot(lep_nu, lep4), r3l), pct(tot(lep_T, lep4), r3l)])
W3 = [1.55, 0.85, 0.7, 1.05, 1.1]

# ---- slide 3: combined totals
T3 = 'Combined sensitivity, all channels'
S3 = 'Hadronic + semileptonic. The flows exist for hadronic only, so their rows use the nu-method semileptonic channels.'
lep_all = list(lep_nu)
old = T({**nu_had, **lep_nu}) / 1.18
S_all = lambda h, l: T({**h, **l})
rows4 = [['method', 'hadronic', 'semileptonic', 'all channels', 'vs old method']]
for m, d in had.items():
    l = lep.get(m, lep_nu)
    rows4.append([m + ('' if m in lep else ' + nu-method τℓ'), f'{T(d):.4f}', f'{T(l):.4f}', f'{S_all(d, l):.4f}', pct(S_all(d, l), old)])
W4 = [3.6, 1.3, 1.5, 1.5, 1.6]
N3 = ['"vs old method": old-method total (0.0529) inferred from the published nu-method +18%; every row uses the same denominator.',
      'Totals are N-weighted: √(Σ N·a² / Σ N). No uncertainty cut: with the common φCP-uncertainty cut (keep 55%) the flows gain about +10 points',
      '(25e +64 circ.: +20.4% → +30.4%; vis+all flow circ.: +26.3% → +36.3%). The transformer has no per-event uncertainty, so no cut.',
      'Semileptonic transformer = vis+all regressor trained on τℓτh only (Oct04).']

NAVY, INK, MUTED, ROW, GOOD, BAD = '1E2761', '212121', '6B7280', 'EEF2F8', '1B6E3A', 'B02A2A'
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

def draw_table(ax, rows, col_w, x0, y0, rh=0.34, size=10, pct_cols=(), bold_last=False):
    for i, row in enumerate(rows):
        y = y0 + i*rh; x = x0
        for j, val in enumerate(row):
            w = col_w[j]; fc = '#'+NAVY if i == 0 else ('#'+ROW if i % 2 == 0 else 'white')
            ax.add_patch(plt.Rectangle((x, y), w, rh, facecolor=fc, edgecolor='white', linewidth=1))
            color = 'white' if i == 0 else '#'+INK; bold = i == 0 or (bold_last and i == len(rows)-1)
            if i > 0 and j in pct_cols and str(val).endswith('%'): color = '#'+BAD if str(val).startswith('−') else '#'+GOOD
            ax.text(x + 0.08 if j == 0 else x + w - 0.08, y + rh/2, str(val), ha='left' if j == 0 else 'right', va='center',
                    fontsize=size, color=color, fontweight='bold' if bold else 'normal')
            x += w

def page(title, subtitle):
    fig = plt.figure(figsize=(13.333, 7.5)); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 13.333); ax.set_ylim(7.5, 0); ax.axis('off')
    ax.text(0.6, 0.65, title, fontsize=24, fontweight='bold', color='#'+NAVY, va='center')
    ax.text(0.6, 1.15, subtitle, fontsize=10.5, color='#'+MUTED, va='center'); return fig, ax

with PdfPages(OUT + '.pdf') as pdf:
    fig, ax = page(T1, S1); draw_table(ax, rows1, W1, 0.6, 1.5, pct_cols=range(3, 8), bold_last=True)
    for i, l in enumerate(N1): ax.text(0.6, 5.75 + 0.3*i, l, fontsize=9.5, color='#'+MUTED)
    pdf.savefig(fig); plt.close(fig)
    fig, ax = page(T2, S2)
    ax.text(0.4, 1.6, 'Hadronic channels with a 3π1π⁰ leg (asymmetry)', fontsize=12, fontweight='bold', color='#'+NAVY)
    draw_table(ax, rows2, W2, 0.4, 1.8)
    ax.text(0.4, 4.25, 'Semileptonic channels (gain over run-3; τℓ – 3π1π⁰ as asymmetry)', fontsize=12, fontweight='bold', color='#'+NAVY)
    draw_table(ax, rows3, W3, 0.4, 4.45, pct_cols=(3, 4), bold_last=True)
    pdf.savefig(fig); plt.close(fig)
    fig, ax = page(T3, S3); draw_table(ax, rows4, W4, 0.6, 1.5, size=11, pct_cols=(4,))
    for i, l in enumerate(N3): ax.text(0.6, 4.0 + 0.3*i, l, fontsize=9.5, color='#'+MUTED)
    pdf.savefig(fig); plt.close(fig)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
rgb = lambda h: RGBColor.from_string(h)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5); blank = prs.slide_layouts[6]

def text(slide, x, y, w, h, s, size=14, bold=False, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run(); r.text = s; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = 'Calibri'

def ptable(slide, x, y, rows, col_w, size=10, pct_cols=(), bold_last=False):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(sum(col_w)), Inches(0.34*len(rows))).table
    for j, cw in enumerate(col_w): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(0.34)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.02); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            r = p.add_run(); r.text = str(val); r.font.size = Pt(size); r.font.name = 'Calibri'; c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = rgb(NAVY); r.font.bold = True; r.font.color.rgb = rgb('FFFFFF')
            else:
                c.fill.fore_color.rgb = rgb(ROW) if i % 2 == 0 else rgb('FFFFFF'); r.font.color.rgb = rgb(INK)
                if j in pct_cols and str(val).endswith('%'): r.font.color.rgb = rgb(BAD) if str(val).startswith('−') else rgb(GOOD)
                if bold_last and i == len(rows)-1: r.font.bold = True

s = prs.slides.add_slide(blank); text(s, 0.6, 0.35, 12.1, 0.6, T1, 28, True, NAVY); text(s, 0.6, 0.95, 12.1, 0.4, S1, 12, False, MUTED)
ptable(s, 0.6, 1.5, rows1, W1, pct_cols=range(3, 8), bold_last=True); text(s, 0.6, 5.7, 12.1, 1.4, ' '.join(N1), 10, False, MUTED)
s = prs.slides.add_slide(blank); text(s, 0.6, 0.35, 12.1, 0.6, T2, 28, True, NAVY); text(s, 0.6, 0.95, 12.1, 0.4, S2, 12, False, MUTED)
text(s, 0.4, 1.45, 8, 0.3, 'Hadronic channels with a 3π1π⁰ leg (asymmetry)', 13, True, NAVY); ptable(s, 0.4, 1.8, rows2, W2)
text(s, 0.4, 4.1, 8, 0.3, 'Semileptonic channels (gain over run-3; τℓ – 3π1π⁰ as asymmetry)', 13, True, NAVY)
ptable(s, 0.4, 4.45, rows3, W3, pct_cols=(3, 4), bold_last=True)
s = prs.slides.add_slide(blank); text(s, 0.6, 0.35, 12.1, 0.6, T3, 28, True, NAVY); text(s, 0.6, 0.95, 12.1, 0.4, S3, 12, False, MUTED)
ptable(s, 0.6, 1.5, rows4, W4, size=11, pct_cols=(4,)); text(s, 0.6, 3.95, 12.1, 1.6, ' '.join(N3), 11, False, MUTED)
prs.save(OUT + '.pptx')
print('written', OUT + '.pdf/.pptx')
for r in rows1 + [[]] + rows2 + [[]] + rows3 + [[]] + rows4:
    print(' | '.join(map(str, r)))
